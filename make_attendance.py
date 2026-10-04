import re, sys
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.page import PageMargins

SRC = sys.argv[1]
OUT = "/home/user/dailywork/Lab_Attendance_Sheets_CSE_AIML.xlsx"
FULL = re.compile(r"^(\d{4}-\d{2}-\d{3})-(\d{3})$")

def expand(spec):
    """'1602-25-733-192, 193, 315 to 317' -> full hall ticket numbers."""
    prefix, out = None, []
    for tok in [t.strip() for t in spec.split(",") if t.strip()]:
        parts = [p.strip() for p in re.split(r"\s+to\s+", tok)]
        nums = []
        for p in parts:
            m = FULL.match(p)
            if m:
                prefix = m.group(1); nums.append(int(m.group(2)))
            else:
                nums.append(int(p))
        lo, hi = nums[0], nums[-1]
        out += [f"{prefix}-{n:03d}" for n in range(lo, hi + 1)]
    return out

src = openpyxl.load_workbook(SRC)["Seating Arrangement"]
labs = {}  # lab -> branch -> list
for r in src.iter_rows(min_row=7, values_only=True):
    if r[0] is None: continue
    lab = re.sub(r"\s+", " ", r[1]).strip()
    br, spec, n = r[2], r[3], r[4]
    nums = expand(spec)
    assert len(nums) == n, (lab, br, len(nums), n)   # cross-check vs sheet count
    labs.setdefault(lab, {})[br] = nums

thin = Side(style="thin"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
bold = Font(name="Calibri", bold=True); hdr_fill = PatternFill("solid", fgColor="D9E1F2")
wb = openpyxl.Workbook(); wb.remove(wb.active)
used = set()

def sheet_name(lab, br):
    s = re.sub(r"[\\/*?:\[\]]", "", lab.split("(")[0]).replace("–", "-").replace(" ", "")
    s = f"{br}_{s}"[:31]
    assert s not in used; used.add(s); return s

for br in ("CSE", "AIML"):
    for lab, d in labs.items():
        if br not in d: continue
        nums = d[br]
        ws = wb.create_sheet(sheet_name(lab, br))
        ws.column_dimensions["A"].width = 7
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 20
        lines = [
            ("VASAVI COLLEGE OF ENGINEERING (AUTONOMOUS), HYDERABAD-31", 14),
            ("TRAINING AND PLACEMENT CELL", 12),
            ("Baseline Assessment (Technical Skills-III) – Attendance Sheet", 12),
            (f"Date: 03-10-2026      Time: 01:20 PM to 02:20 PM", 11),
            (f"Lab: {lab}", 11),
            (f"Branch: {'CSE-AIML' if br == 'AIML' else 'CSE'}      Students Allotted: {len(nums)}", 11),
        ]
        for i, (t, sz) in enumerate(lines, 1):
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
            c = ws.cell(i, 1, t); c.font = Font(name="Calibri", bold=True, size=sz)
            c.alignment = Alignment(horizontal="center" if i <= 4 else "left")
        h = 8
        for j, t in enumerate(["S. No.", "Hall Ticket No.", "Signature", "Remarks"], 1):
            c = ws.cell(h, j, t); c.font = bold; c.fill = hdr_fill; c.border = box
            c.alignment = Alignment(horizontal="center", vertical="center")
        for i, hn in enumerate(nums, 1):
            r = h + i
            ws.row_dimensions[r].height = 24
            for j, v in enumerate([i, hn, None, None], 1):
                c = ws.cell(r, j, v); c.border = box
                c.alignment = Alignment(horizontal="center" if j <= 2 else "left", vertical="center")
        end = h + len(nums)
        sr = end + 3
        ws.cell(sr, 1, "Present: ______    Absent: ______"); 
        ws.print_title_rows = f"{h}:{h}"
        ws.page_setup.orientation = "portrait"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
        ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.5, bottom=0.5)
        ws.print_area = f"A1:D{sr}"
wb.save(OUT)
tot = {b: sum(len(d.get(b, [])) for d in labs.values()) for b in ("CSE", "AIML")}
print(len(wb.sheetnames), "sheets;", tot, wb.sheetnames)
